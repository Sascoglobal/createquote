"""Extend the user's Template_design.xlsx with app-control columns (keeps their layout and values)."""
import re, copy, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
wb = load_workbook(SRC)

F = "Calibri"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
fill_blue = PatternFill("solid", fgColor="DDEBF7")   # decision columns pre-set by a draft rule
fill_yel = PatternFill("solid", fgColor="FFFF99")    # you fill in
fill_prop = PatternFill("solid", fgColor="FCE4D6")   # proposed by Claude, confirm or delete
hdr_fill = PatternFill("solid", fgColor="1F4E79")
hdr_font = Font(name=F, size=11, bold=True, color="FFFFFF")

APP_HEADERS = ["Direction", "Load Type", "Cargo Nature", "Cargo Type", "Incoterms",
               "Default On", "VAT Applies", "Rate Control", "Margin %", "Default Qty"]
APP_COLS = list(range(8, 8 + len(APP_HEADERS)))      # H..Q
WIDTHS = [11, 11, 13, 14, 14, 11, 11, 13, 10, 11]
ID_RE = re.compile(r"^[A-Z]{2}-[A-Z]{3}-\d{2}")
INCOS = ["EXW", "FOB", "FCA", "FAS", "CFR", "CIF", "CPT", "CIP", "DAP", "DPU", "DDP"]

SHEETS = {  # sheet -> (mode, direction)
    "Sea Import": ("Sea", "Import"), "Sea Export": ("Sea", "Export"),
    "Air Import": ("Air", "Import"), "Air Export": ("Air", "Export"),
    "Land Transport": ("Land", "ALL"),
}

CORE = ["thc", "terminal handling", "bill of lading", "bl issuance", "bill of lading / telex", "documentation", "delivery order",
        "do fee", "customs clearance", "customs declaration", "awb", "pick-up & loading", "pickup", "pick-up",
        "jebel ali to dubai", "airport to warehouse", "last mile", "destination final mile", "door delivery", "border export customs"]
AUTO = ["thc", "terminal handling", "bill of lading", "bl issuance", "documentation", "delivery order", "do fee", "customs clearance",
        "customs declaration", "awb", "dpc", "tluc", "vgm", "token", "mofaic", "ministry", "importer on record", "label", "mcc", "manifest"]
SURCH = ["surcharge", "bunker", "baf", "war risk", "peak", "fuel", "security", "toll", "border crossing", "waiting", "detention", "storage", "demurrage", "inspection"]


def classify(desc, unit, rem, cat_text, sheet):
    d = (desc or "").strip(); u = (unit or ""); r = (rem or "")
    text = (d + " " + r).lower()
    pct = "%" in u
    passthrough = pct and bool(re.search(r"duty|duties|\bvat\b|\btax", text))
    cat = cat_text.lower()
    freight = cat.startswith("freight")
    variant = bool(re.search(r"20ft|40hc|40ft|20'|40'|per cbm\(for lcl\)", text)) and "ocean freight per container" not in text
    on = "N"
    if not passthrough:
        if freight:
            if not variant and not any(k in text for k in SURCH) and not re.search(r"\bdg\b|imo|dangerous", text):
                on = "Y"
        elif any(k in text for k in CORE) and not any(k in text for k in ("inspection", "waiting", "detention", "storage", "demurrage")):
            on = "Y"
    rc = "Auto" if (not passthrough and not freight and any(k in text for k in AUTO) and not pct) else "Editable"
    vat = "N" if passthrough else "Y"
    ctype = "ALL"
    if re.search(r"\bdg\b|imo|dangerous", text): ctype = "Dangerous"
    elif re.search(r"refrigerat|reefer", text): ctype = "Refrigerated"
    codes = re.findall(r"\b(EXW|FOB|FCA|FAS|CFR|CIF|CPT|CIP|DAP|DPU|DDP)\b", (d + " " + r).upper())
    inco = ",".join(dict.fromkeys(codes)) if codes else "ALL"
    load = "ALL"
    if sheet.startswith("Sea") and "container" in u.lower(): load = "FCL"
    if sheet == "Land Transport" and re.search(r"13\.6m|3 ton pickup", u.lower()): load = "FTL"
    return dict(on=on, rc=rc, vat=vat, ctype=ctype, inco=inco, load=load, pct=pct, passthrough=passthrough)


def style_hdr(ws, row, headers, cols):
    for h, c in zip(headers, cols):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill; cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, cols, widths):
    for c, w in zip(cols, widths):
        ws.column_dimensions[get_column_letter(c)].width = w


def dv(ws, formula, rng):
    d = DataValidation(type="list", formula1='"%s"' % formula, allow_blank=True)
    ws.add_data_validation(d); d.add(rng)


report = {"lines": 0, "blank_unit": [], "blank_ccy": [], "dup_desc": [], "pct_fmt": 0, "by_sheet": {}}
for sheet, (mode, direction) in SHEETS.items():
    ws = wb[sheet]
    seen = {}
    sec = ""
    n = 0
    for row in range(1, ws.max_row + 1):
        a = ws.cell(row=row, column=1).value
        if isinstance(a, str) and a.startswith("SECTION:"):
            sec = a.replace("SECTION:", "").strip()
        elif a == "Cost Item ID":
            style_hdr(ws, row, APP_HEADERS, APP_COLS)
        elif isinstance(a, str) and ID_RE.match(a):
            desc = ws.cell(row=row, column=2).value
            cat = ws.cell(row=row, column=3).value or sec.title()
            unit = ws.cell(row=row, column=4).value
            ccy = ws.cell(row=row, column=5).value
            rem = ws.cell(row=row, column=7).value
            k = classify(desc, unit, rem, cat, sheet)
            vals = [direction, k["load"], "ALL", k["ctype"], k["inco"], k["on"], k["vat"], k["rc"], None, 1]
            for col, v in zip(APP_COLS, vals):
                c = ws.cell(row=row, column=col, value=v)
                c.font = Font(name=F, size=11); c.border = border; c.alignment = Alignment(horizontal="center")
                c.fill = fill_yel if col == 16 else fill_blue
            ws.cell(row=row, column=16).number_format = "0.0%"
            fc = ws.cell(row=row, column=6)
            fc.number_format = "0.00%" if k["pct"] else "#,##0.00"
            if k["pct"]: report["pct_fmt"] += 1
            n += 1
            if not unit: report["blank_unit"].append(a)
            if not ccy: report["blank_ccy"].append(a)
            key = (str(desc or "").strip().lower())
            if key in seen: report["dup_desc"].append((seen[key], a, desc))
            seen[key] = a
    set_widths(ws, APP_COLS, WIDTHS)
    ws.cell(row=1, column=9, value="Columns H to Q control how the quote app uses each line. See the 'App Guide' sheet.").font = Font(name=F, size=10, italic=True, color="7F7F7F")
    last = ws.max_row + 150
    dv(ws, "ALL,Import,Export", "H5:H%d" % last)
    dv(ws, "ALL,FCL,LCL,FTL,LTL", "I5:I%d" % last)
    dv(ws, "ALL,Palletized,Loose,Mixed", "J5:J%d" % last)
    dv(ws, "ALL,General,Refrigerated,OOG,Dangerous,Other", "K5:K%d" % last)
    dv(ws, "Y,N", "M5:N%d" % last)
    dv(ws, "Auto,Editable", "O5:O%d" % last)
    report["by_sheet"][sheet] = n; report["lines"] += n

# ---------------------------------------------------------------- Warehousing
ws = wb["Warehousing Rates"]
WH_HEAD = ["Cargo Type", "Default On", "VAT Applies", "Rate Control", "Margin %", "Default Qty"]
WH_COLS = list(range(8, 8 + len(WH_HEAD)))
WH_MAP = {"general": "General", "ambient": "Refrigerated", "cold": "Refrigerated", "dangerous": "Dangerous", "excise": "Other"}
wh_n = 0
for row in range(1, ws.max_row + 1):
    a = ws.cell(row=row, column=1).value
    if a == "Storage Category":
        style_hdr(ws, row, WH_HEAD, WH_COLS)
    elif isinstance(a, str) and ws.cell(row=row, column=2).value and not a.startswith("CATEGORY") and a != "Storage Category" and ws.cell(row=row, column=3).value:
        key = a.lower()
        ctype = next((v for k, v in WH_MAP.items() if key.startswith(k)), "ALL")
        act = str(ws.cell(row=row, column=2).value).lower()
        on = "Y" if ("storage" in act and "audit" not in act) or "in/out" in act or "handling" in act else "N"
        vals = [ctype, on, "Y", "Editable", None, 1]
        for col, v in zip(WH_COLS, vals):
            c = ws.cell(row=row, column=col, value=v)
            c.font = Font(name=F, size=11); c.border = border; c.alignment = Alignment(horizontal="center")
            c.fill = fill_yel if col == 12 else fill_blue
        ws.cell(row=row, column=12).number_format = "0.0%"
        ws.cell(row=row, column=5).number_format = "#,##0.00"
        wh_n += 1
set_widths(ws, WH_COLS, [14, 11, 11, 13, 10, 11])
ws.cell(row=1, column=9, value="Columns H to M control how the quote app uses each line. See the 'App Guide' sheet.").font = Font(name=F, size=10, italic=True, color="7F7F7F")
dv(ws, "ALL,General,Refrigerated,OOG,Dangerous,Other", "H5:H%d" % (ws.max_row + 150))
dv(ws, "Y,N", "I5:J%d" % (ws.max_row + 150))
dv(ws, "Auto,Editable", "K5:K%d" % (ws.max_row + 150))
report["by_sheet"]["Warehousing Rates"] = wh_n; report["lines"] += wh_n

# ---------------------------------------------------------------- Add-ons (proposed by Claude)
wa = wb.create_sheet("Add-ons (proposed)")
wa["A1"] = "ADD-ONS PROPOSED BY CLAUDE (cargo type, LCL and LTL lines). Review each row: fix, fill the cost, or delete the row."
wa["A1"].font = Font(name=F, size=12, bold=True, color="1F4E79")
wa["A2"] = "Same columns as the pricing sheets, plus Mode in column R. These lines were not in your Template_design; they fill gaps for Refrigerated, OOG, Dangerous, Other cargo, LCL and LTL."
wa["A2"].font = Font(name=F, size=10, italic=True)
heads = ["Cost Item ID", "Cost Description", "Cost Category", "Basis / Unit", "Currency", "Base Cost Rate", "Remarks / Notes"] + APP_HEADERS + ["Mode", "Status"]
for i, h in enumerate(heads, 1):
    c = wa.cell(row=4, column=i, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
# (mode, dir, load, nature, ctype, cat, desc, unit, ccy, on, rc, remarks)
ADD = [
    ("Sea", "ALL", "FCL", "ALL", "Refrigerated", "Freight Charges", "Reefer Surcharge", "Per Container", "USD", "Y", "Editable", "Carrier tariff"),
    ("Sea", "ALL", "FCL", "ALL", "Refrigerated", "Destination Charges", "Reefer Plug-in / Monitoring", "Per Container / Day", "AED", "N", "Editable", "At actuals"),
    ("Sea", "ALL", "FCL", "ALL", "Refrigerated", "Origin Charges", "Pre-Trip Inspection (PTI)", "Per Container", "AED", "N", "Auto", ""),
    ("Sea", "ALL", "FCL", "ALL", "Refrigerated", "Delivery Charges", "Genset / Reefer Trucking Premium", "Per Container", "AED", "N", "Editable", "If trucked"),
    ("Air", "ALL", "ALL", "ALL", "Refrigerated", "Freight Charges", "Temperature-Controlled Handling Surcharge", "Per KG", "USD", "Y", "Editable", "Airline approval required"),
    ("Air", "ALL", "ALL", "ALL", "Refrigerated", "Destination Charges", "Cool Room Storage", "Per KG / Day", "AED", "N", "Editable", "At actuals"),
    ("Land", "ALL", "ALL", "ALL", "Refrigerated", "Freight Charges", "Reefer Truck Premium", "Per Trip", "AED", "Y", "Editable", ""),
    ("Sea", "ALL", "ALL", "ALL", "OOG", "Freight Charges", "OOG / Special Equipment Surcharge (FR / OT)", "Per Container", "USD", "Y", "Editable", "Carrier acceptance required"),
    ("Sea", "ALL", "ALL", "ALL", "OOG", "Origin Charges", "Lashing, Securing & Survey", "Per Shipment", "AED", "N", "Editable", ""),
    ("Sea", "ALL", "ALL", "ALL", "OOG", "Destination Charges", "Port Permits / Special Handling", "Per Shipment", "AED", "N", "Editable", "At actuals"),
    ("Air", "ALL", "ALL", "ALL", "OOG", "Freight Charges", "Outsize Cargo / Charter Handling", "Per Shipment", "USD", "Y", "Editable", "Quote on request"),
    ("Land", "ALL", "ALL", "ALL", "OOG", "Freight Charges", "Low-Bed Trailer Premium", "Per Trip", "AED", "Y", "Editable", ""),
    ("Land", "ALL", "ALL", "ALL", "OOG", "Freight Charges", "Special Permits & Escort", "Per Trip", "AED", "N", "Editable", "At actuals"),
    ("Sea", "Import", "ALL", "ALL", "Dangerous", "Freight Charges", "DG Surcharge", "Per Container", "USD", "Y", "Editable", "Export already has IMO DG Surcharge in your template"),
    ("Sea", "ALL", "ALL", "ALL", "Dangerous", "Origin Charges", "DG Declaration / IMDG Documentation", "Per Shipment", "AED", "Y", "Auto", ""),
    ("Sea", "ALL", "ALL", "ALL", "Dangerous", "Destination Charges", "DG Port Handling / Storage", "Per Container / Day", "AED", "N", "Editable", "At actuals"),
    ("Air", "ALL", "ALL", "ALL", "Dangerous", "Origin Charges", "DG Acceptance Check / Handling (IATA DGR)", "Per Shipment", "USD", "Y", "Auto", ""),
    ("Air", "Import", "ALL", "ALL", "Dangerous", "Freight Charges", "DG Surcharge", "Per KG", "USD", "Y", "Editable", "Export already has DG Surcharges in your template"),
    ("Land", "ALL", "ALL", "ALL", "Dangerous", "Freight Charges", "DG Vehicle Premium / Permit", "Per Trip", "AED", "Y", "Editable", ""),
    ("Sea,Air,Land", "ALL", "ALL", "ALL", "Other", "Other Charges", "Special Cargo Handling (specify)", "Per Shipment", "AED", "Y", "Editable", "Describe in remarks"),
    ("Sea,Air,Land", "ALL", "ALL", "Palletized", "ALL", "Origin Charges", "Pallet Handling / Shrink Wrapping", "Per Pallet", "AED", "N", "Editable", ""),
    ("Sea,Air,Land", "ALL", "ALL", "Loose", "ALL", "Origin Charges", "Loose Cargo Manual Handling", "Per CBM", "AED", "Y", "Editable", "Floor-loaded / manual unloading"),
    ("Sea,Air,Land", "ALL", "ALL", "Mixed", "ALL", "Origin Charges", "Segregation / Sorting / Labelling", "Per Shipment", "AED", "Y", "Editable", ""),
    ("Sea", "Import", "LCL", "ALL", "ALL", "Origin Charges", "Origin CFS / Consolidation Charges", "Per CBM", "USD", "Y", "Editable", "Min 1 CBM"),
    ("Sea", "Import", "LCL", "ALL", "ALL", "Freight Charges", "Ocean Freight (LCL)", "Per CBM / W/M", "USD", "Y", "Editable", "Higher of weight or measure"),
    ("Sea", "Import", "LCL", "ALL", "ALL", "Destination Charges", "Destination CFS / De-consolidation Charges", "Per CBM", "AED", "Y", "Editable", "At Jebel Ali CFS"),
    ("Sea", "Export", "LCL", "ALL", "ALL", "Origin Charges", "Origin CFS / Stuffing Charges", "Per CBM", "AED", "Y", "Editable", "Min 1 CBM"),
    ("Sea", "Export", "LCL", "ALL", "ALL", "Freight Charges", "Ocean Freight (LCL)", "Per CBM / W/M", "USD", "Y", "Editable", "Higher of weight or measure"),
    ("Land", "ALL", "LTL", "ALL", "ALL", "Freight Charges", "Inter-GCC Freight: LTL (part load)", "Per CBM", "AED", "Y", "Editable", ""),
]
for i, (mode, d, load, nat, ct, cat, desc, unit, ccy, on, rc, rem) in enumerate(ADD, 5):
    vals = ["AD-%03d" % (i - 4), desc, cat, unit, ccy, None, rem, d, load, nat, ct, "ALL", on, "Y", rc, None, 1, mode, "Proposed - confirm"]
    for j, v in enumerate(vals, 1):
        c = wa.cell(row=i, column=j, value=v)
        c.font = Font(name=F, size=11); c.border = border
        c.alignment = Alignment(horizontal="left" if j in (2, 3, 7) else "center", wrap_text=(j in (2, 7)))
        if j in (8, 9, 10, 11, 12, 13, 14, 15, 17, 18): c.fill = fill_blue
    wa.cell(row=i, column=6).fill = fill_yel; wa.cell(row=i, column=16).fill = fill_yel
    wa.cell(row=i, column=19).fill = fill_prop
    wa.cell(row=i, column=6).number_format = "#,##0.00"; wa.cell(row=i, column=16).number_format = "0.0%"
for col, w in zip("ABCDEFGHIJKLMNOPQRS", [12, 44, 20, 20, 9, 14, 34, 11, 11, 13, 14, 12, 11, 11, 13, 10, 11, 14, 20]):
    wa.column_dimensions[col].width = w
wa.freeze_panes = "C5"
last = 5 + len(ADD) + 100
dv(wa, "ALL,Import,Export", "H5:H%d" % last); dv(wa, "ALL,FCL,LCL,FTL,LTL", "I5:I%d" % last)
dv(wa, "ALL,Palletized,Loose,Mixed", "J5:J%d" % last); dv(wa, "ALL,General,Refrigerated,OOG,Dangerous,Other", "K5:K%d" % last)
dv(wa, "Y,N", "M5:N%d" % last); dv(wa, "Auto,Editable", "O5:O%d" % last)
report["by_sheet"]["Add-ons (proposed)"] = len(ADD)

# ---------------------------------------------------------------- App Guide
wg = wb.create_sheet("App Guide")
wg["A1"] = "How the quote app reads this workbook"; wg["A1"].font = Font(name=F, size=14, bold=True, color="1F4E79")
guide = [
    ("Your layout is unchanged", "Every sheet, value and formula you made is still here. I only added columns on the right (blue and yellow) and two sheets: 'Add-ons (proposed)' and this guide."),
    ("Base Cost Rate (col F)", "Your cost for the line, in the Currency in column E. Leave blank until you have the number. For lines whose unit is a % (insurance, duty), enter the percentage as a % (0.35% is typed 0.35% or 0.0035)."),
    ("Margin % (yellow)", "Optional. Sell rate = Base Cost Rate x (1 + Margin %). Type 15% or 0.15. The salesperson sees only the sell rate, never the cost or the margin."),
    ("Direction / Load Type / Cargo Nature / Cargo Type", "ALL = shows for every choice. Otherwise the line shows only when the quote matches. Load Type is FCL, LCL, FTL or LTL. Cargo Type is General, Refrigerated, OOG, Dangerous or Other."),
    ("Incoterms", "ALL, or a list such as EXW or DDP,CFR. When a list is given, the line shows only for those Incoterms and skips the Incoterm_Matrix rule. When ALL, the Incoterm_Matrix decides by category."),
    ("Default On", "Y = the line is ticked (included) when the quote opens. N = it is there but unticked."),
    ("VAT Applies", "Y = the line is VAT-chargeable by default. The salesperson can tick or untick VAT per line on the quote."),
    ("Rate Control", "Auto = once a cost is filled in, the rate is locked and cannot be edited on the quote. Editable = the salesperson can change the rate."),
    ("Default Qty", "Usually 1. For % lines the quantity on the quote is the cargo value (CIF value) in the quote currency."),
    ("Category order on the quote", "Freight, Origin, Destination, Customs, Delivery, Insurance. Warehousing lines follow their own storage categories."),
    ("Not wired yet", "The 'Emirates & Truck Matrix' and 'Quote Calculator' sheets are not read by the app yet."),
    ("Colours", "Blue = pre-set by a draft rule, change it freely. Yellow = you fill in. Orange (Add-ons sheet) = proposed by Claude, confirm or delete."),
]
for i, (a, b) in enumerate(guide, 3):
    wg.cell(row=i, column=1, value=a).font = Font(name=F, size=11, bold=True)
    c = wg.cell(row=i, column=2, value=b); c.font = Font(name=F, size=11); c.alignment = Alignment(wrap_text=True, vertical="top")
    wg.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
wg.column_dimensions["A"].width = 38; wg.column_dimensions["B"].width = 120

wb.save(OUT)
import json
print(json.dumps({k: (v if k != "dup_desc" else v) for k, v in report.items()}, indent=1, default=str))
