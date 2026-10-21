"""Reads Sasco_Quote_Templates_v3.xlsx (settings) and Template_design_v2.xlsx (line items) -> data.json for the quote app.
usage: python3 xlsx_to_data.py <templates.xlsx> <design.xlsx> <out.json>"""
import json, re, sys
from openpyxl import load_workbook

TPL, DES, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
issues = []


def s(v):
    return "" if v is None else str(v).strip()


def num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def frac(v, what):
    """fraction (0.15) expected; a value >1 is read as a percent number (15 -> 0.15)."""
    n = num(v)
    if n is None: return None
    if n > 1:
        issues.append("%s: %s read as %s%%" % (what, v, v)); return n / 100.0
    return n


wt = load_workbook(TPL, data_only=True)

# ---------------- settings
keymap = {
    "Company name": "company_name", "Company email": "company_email", "Address (printed in PDF header)": "address",
    "Phone (printed in PDF header)": "phone", "Website": "website", "Company TRN (VAT registration no.)": "trn",
    "Quote number prefix": "quote_prefix", "Default validity (days)": "validity_days", "Default currency": "currency",
    "Default VAT rate": "vat_rate", "Default payment terms line": "payment_terms", "Email subject format": "email_subject",
    "Email greeting": "email_greeting", "Email opening line": "email_open", "Email closing line": "email_close",
    "Email signature": "email_sig", "Bank details printed on quote? (Y/N)": "bank_print", "Bank - beneficiary / bank / IBAN / SWIFT": "bank_details",
}
settings = {}
for r in wt["Settings"].iter_rows(min_row=4, values_only=True):
    if r[0] in keymap:
        settings[keymap[r[0]]] = r[1] if r[1] is not None else ""
vd = num(settings.get("validity_days"))
settings["validity_days"] = int(vd) if vd else None
if vd is None and s(settings.get("validity_days")) == "":
    pass
vr = num(settings.get("vat_rate"))
settings["vat_rate"] = vr if vr is not None else 0.05
if s(settings.get("currency")) == "": settings["currency"] = "AED"
DEF_CCY = s(settings["currency"]).upper()

# ---------------- FX
rates, names = {}, {}
fee = 0.0
for r in wt["FX"].iter_rows(min_row=5, values_only=True):
    if s(r[0]) and len(s(r[0])) == 3 and num(r[2]) is not None and num(r[2]) > 0:
        rates[s(r[0]).upper()] = num(r[2]); names[s(r[0]).upper()] = s(r[1])
    elif s(r[0]) and len(s(r[0])) == 3:
        names[s(r[0]).upper()] = s(r[1]); rates.setdefault(s(r[0]).upper(), None)
    if s(r[0]).lower().startswith("currency conversion"):
        f = num(r[2]); fee = 0.0 if f is None else (f / 100.0 if f >= 1 else f)
rates["AED"] = 1.0
fx = {"rates": rates, "names": names, "fee": fee}
for c, v in rates.items():
    if v is None: issues.append("FX: no AED rate for %s yet" % c)

# ---------------- team
team = []
ws = wt["Sales_Team"]
hdr = {s(c.value).lower(): i for i, c in enumerate(ws[1])}
for r in ws.iter_rows(min_row=2, values_only=True):
    email = s(r[hdr.get("email (@sascoglobal.com)", 1)]).lower()
    if not email or "@" not in email: continue
    if not re.match(r"^[^@\s]+@sascoglobal\.com$", email): issues.append("Team: email looks wrong: %s" % email)
    perms = s(r[hdr["permissions"]]) if "permissions" in hdr else ""
    act = s(r[hdr["active (y/n)"]]) if "active (y/n)" in hdr else "Y"
    team.append({"name": s(r[hdr.get("name", 0)]), "email": email, "role": s(r[hdr.get("role", 2)]),
                 "dash": "dashboard" in perms.lower(), "active": act.upper() != "N"})

# ---------------- matrix
matrix = {}
cats6 = ["ORIGIN", "FREIGHT", "INSURANCE", "DESTINATION", "CUSTOMS", "DELIVERY"]
for r in wt["Incoterm_Matrix"].iter_rows(min_row=5, values_only=True):
    if not r[0] or r[1] not in ("Import", "Export"): continue
    matrix.setdefault(s(r[0]), {})[s(r[1])] = {c: (s(r[3 + i]) or "Y") for i, c in enumerate(cats6)}

# ---------------- terms
terms = []
for r in wt["Terms"].iter_rows(min_row=2, values_only=True):
    if not r[0] or not s(r[4]): continue
    terms.append({"id": s(r[0]), "layer": s(r[1]), "key": s(r[2]) or "ALL", "mode": s(r[3]) or "ALL",
                  "text": s(r[4]), "on": s(r[5]) != "N", "untick": s(r[6]) != "N"})

# ---------------- charges
wd = load_workbook(DES, data_only=True)
ID_RE = re.compile(r"^[A-Z]{2}-[A-Z]{3}-\d{2}")
SHEETS = {"Sea Import": ("Sea", "Import"), "Sea Export": ("Sea", "Export"), "Air Import": ("Air", "Import"),
          "Air Export": ("Air", "Export"), "Land Transport": ("Land", "ALL")}


def cat_of(text):
    t = s(text).lower()
    for k, v in (("origin", "ORIGIN"), ("freight", "FREIGHT"), ("destination", "DESTINATION"), ("customs", "CUSTOMS"),
                 ("delivery", "DELIVERY"), ("transportation", "DELIVERY"), ("insurance", "INSURANCE")):
        if t.startswith(k): return v
    return "OTHER"


def yn(v, d):
    v = s(v).upper()
    return v if v in ("Y", "N") else d


charges = []


def add(sheet_mode, row, section_default):
    d = row
    unit = d["unit"]
    pct = "%" in unit
    cost = d["cost"]
    if pct and cost is not None and cost > 1:
        issues.append("%s: %% cost %s read as %s%%" % (d["id"], cost, cost)); cost = cost / 100.0
    margin = frac(d.get("margin"), "%s margin" % d["id"])
    ccy = s(d["ccy"]).upper() or DEF_CCY
    qty = num(d.get("qty")); qty = 1 if qty is None else qty
    charges.append({
        "id": d["id"], "mode": d["mode"], "dir": d["dir"] or "ALL", "load": d["load"] or "ALL", "nature": d["nature"] or "ALL",
        "ctype": d["ctype"] or "ALL", "inco": d["inco"] or "ALL", "cat": d["cat"], "sec": d["sec"], "d": d["desc"],
        "u": unit or "Per Shipment", "ccy": ccy, "cost": cost, "margin": margin, "q": qty,
        "rc": d["rc"] if d["rc"] in ("Auto", "Editable") else "Editable", "on": yn(d["on"], "N") == "Y", "vat": yn(d["vat"], "Y") == "Y",
        "rem": d["rem"], "pct": pct,
    })


for sheet, (mode, ddir) in SHEETS.items():
    ws = wd[sheet]
    sec = ""; hmap = {}
    for r in ws.iter_rows(min_row=1, values_only=False):
        a = r[0].value
        if isinstance(a, str) and a.startswith("SECTION:"):
            sec = a.replace("SECTION:", "").strip()
        elif a == "Cost Item ID":
            hmap = {s(c.value).lower(): i for i, c in enumerate(r) if c.value}
        elif isinstance(a, str) and ID_RE.match(a):
            g = lambda k, default=None: r[hmap[k]].value if k in hmap and hmap[k] < len(r) else default
            catt = cat_of(g("cost category") or sec)
            add(mode, {"id": s(a), "mode": mode, "dir": s(g("direction")) or ddir, "load": s(g("load type")), "nature": s(g("cargo nature")),
                       "ctype": s(g("cargo type")), "inco": s(g("incoterms")), "cat": catt,
                       "sec": catt, "desc": s(g("cost description")), "unit": s(g("basis / unit")), "ccy": g("currency"),
                       "cost": num(g("base cost rate")), "margin": g("margin %"), "qty": g("default qty"), "rc": s(g("rate control")),
                       "on": g("default on"), "vat": g("vat applies"), "rem": s(g("remarks / notes"))}, sec)
            if not s(g("basis / unit")): issues.append("%s: no unit" % a)
            if not s(g("currency")): issues.append("%s: no currency (using %s)" % (a, DEF_CCY))

# warehousing
ws = wd["Warehousing Rates"]
hmap = {}; n = 0
for r in ws.iter_rows(min_row=1, values_only=False):
    a = r[0].value
    if a == "Storage Category":
        hmap = {s(c.value).lower(): i for i, c in enumerate(r) if c.value}
    elif isinstance(a, str) and not a.startswith("CATEGORY") and a != "Storage Category" and len(r) > 2 and r[1].value and r[2].value and hmap:
        g = lambda k, default=None: r[hmap[k]].value if k in hmap and hmap[k] < len(r) else default
        n += 1
        desc = s(r[1].value); temp = s(g("temp range"))
        if "°" in temp: desc += " (%s)" % temp
        cost_key = next((k for k in hmap if k.startswith("base cost rate")), None)
        add("Warehouse", {"id": "WH-%02d" % n, "mode": "Warehouse", "dir": "ALL", "load": "ALL", "nature": "ALL", "ctype": s(g("cargo type")),
                          "inco": "ALL", "cat": "WAREHOUSE", "sec": s(a).upper(), "desc": desc, "unit": s(g("billing basis / unit")),
                          "ccy": "AED", "cost": num(r[hmap[cost_key]].value) if cost_key else None, "margin": g("margin %"), "qty": g("default qty"),
                          "rc": s(g("rate control")), "on": g("default on"), "vat": g("vat applies"), "rem": s(g("remarks"))}, "")

# add-ons
if "Add-ons (proposed)" in wd.sheetnames:
    ws = wd["Add-ons (proposed)"]
    hmap = {}
    for r in ws.iter_rows(min_row=1, values_only=False):
        a = r[0].value
        if a == "Cost Item ID":
            hmap = {s(c.value).lower(): i for i, c in enumerate(r) if c.value}
        elif isinstance(a, str) and a.startswith("AD-") and hmap:
            g = lambda k, default=None: r[hmap[k]].value if k in hmap and hmap[k] < len(r) else default
            catt = cat_of(g("cost category"))
            add(s(g("mode")), {"id": s(a), "mode": s(g("mode")) or "ALL", "dir": s(g("direction")), "load": s(g("load type")), "nature": s(g("cargo nature")),
                               "ctype": s(g("cargo type")), "inco": s(g("incoterms")), "cat": catt, "sec": catt, "desc": s(g("cost description")),
                               "unit": s(g("basis / unit")), "ccy": g("currency"), "cost": num(g("base cost rate")), "margin": g("margin %"),
                               "qty": g("default qty"), "rc": s(g("rate control")), "on": g("default on"), "vat": g("vat applies"), "rem": s(g("remarks / notes"))}, "")

filled = sum(1 for c in charges if c["cost"] is not None)
json.dump({"charges": charges, "terms": terms, "matrix": matrix, "settings": settings, "fx": fx, "team": team},
          open(OUT, "w"), ensure_ascii=False)
print("%d charges (%d with a cost), %d terms, %d incoterms, %d team, %d currencies" % (len(charges), filled, len(terms), len(matrix), len(team), len(rates)))
for i in issues: print(" - ", i)
