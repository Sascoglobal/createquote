"""Update the user's Sasco_Quote_Templates.xlsx: keep their edits, add FX, fix matrix insurance, drop obsolete Charges sheet."""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]
wb = load_workbook(SRC)
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="1F3864"); hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF")
yel = PatternFill("solid", fgColor="FFFF00"); org = PatternFill("solid", fgColor="F8CBAD")
changes = []

# 1) obsolete sheets
for name in ("Charges", "Sheet2"):
    if name in wb.sheetnames:
        del wb[name]

# 2) matrix: insurance column (F). Import quotes the buyer's side; buyer insures under EXW/FCA/FAS/FOB/CFR/CPT.
ws = wb["Incoterm_Matrix"]
buyer_insures = {"EXW", "FCA", "FAS", "FOB", "CFR", "CPT"}
seller_insures = {"CIF", "CIP", "DAP", "DPU", "DDP"}
for r in range(5, 27):
    inc, d = ws.cell(r, 1).value, ws.cell(r, 2).value
    if d == "Import": want = "Y" if inc in buyer_insures else "N"
    elif d == "Export": want = "Y" if inc in seller_insures else "N"
    else: continue
    c = ws.cell(r, 6)
    if c.value != want:
        changes.append((d, inc, c.value, want))
        c.value = want; c.fill = org
        ws.cell(r, 10).value = "Insurance changed by Claude (was %s). Buyer insures under %s; seller insures under CIF/CIP and DAP/DPU/DDP." % (changes[-1][2], "EXW/FCA/FAS/FOB/CFR/CPT")
        ws.cell(r, 10).font = Font(name=F, size=9, italic=True)
ws["A2"].value = (str(ws["A2"].value or "") + " Orange cells were changed by Claude so that the Insurance lines in your Template_design can appear; change them back if you disagree.").strip()

# 3) sales team: fix truncated email
st = wb["Sales_Team"]
fixed_email = None
for r in range(2, st.max_row + 1):
    v = st.cell(r, 2).value
    if isinstance(v, str) and v.endswith("@sascoglobal."):
        fixed_email = (v, v + "com"); st.cell(r, 2).value = v + "com"; st.cell(r, 2).fill = org

# 4) FX sheet
fx = wb.create_sheet("FX")
fx["A1"] = "Currencies and conversion"; fx["A1"].font = Font(name=F, size=16, bold=True, color="1F3864")
fx["A2"] = ("Line costs sit in the currency shown in Template_design. The salesperson picks the quote currency on each quote. "
            "Amounts are converted through AED using the rates below, plus the transfer charge. Only the final converted amounts are printed.")
fx["A2"].font = Font(name=F, size=9, italic=True); fx["A2"].alignment = Alignment(wrap_text=True, vertical="top")
fx.merge_cells("A2:D2"); fx.row_dimensions[2].height = 40
for i, h in enumerate(["Currency", "Name", "AED per 1 unit", "Note"], 1):
    c = fx.cell(row=4, column=i, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
rows = [
    ("AED", "UAE Dirham", 1, "Base currency, always 1"),
    ("USD", "US Dollar", 3.6725, "Official AED/USD peg. Change only if the peg changes."),
    ("EUR", "Euro", None, "Fill in your rate. Market-driven, so it is not pre-filled."),
    ("GBP", "British Pound", None, "Fill in your rate. Market-driven, so it is not pre-filled."),
]
for i, row in enumerate(rows, 5):
    for j, v in enumerate(row, 1):
        c = fx.cell(row=i, column=j, value=v); c.font = Font(name=F, size=10, color="0000FF" if j == 3 else "000000"); c.border = border
    fx.cell(row=i, column=3).number_format = "0.0000"
    if row[2] is None: fx.cell(row=i, column=3).fill = yel
for i in range(9, 14):  # spare rows for more currencies
    for j in range(1, 5):
        c = fx.cell(row=i, column=j); c.border = border; c.font = Font(name=F, size=10, color="0000FF")
    fx.cell(row=i, column=3).number_format = "0.0000"; fx.cell(row=i, column=1).fill = yel
fx["A15"] = "Currency conversion / transfer charge (%)"; fx["A15"].font = Font(name=F, size=10, bold=True)
fx["C15"] = 0; fx["C15"].number_format = "0.00%"; fx["C15"].fill = yel; fx["C15"].font = Font(name=F, size=10, color="0000FF"); fx["C15"].border = border
fx["D15"] = "Added on top of a converted line when its cost currency differs from the quote currency. 0% = none. Salespeople can override per quote."
fx["D15"].font = Font(name=F, size=9, italic=True)
for col, w in zip("ABCD", [34, 20, 18, 90]): fx.column_dimensions[col].width = w
fx["A17"] = "Add more currencies in the yellow spare rows (for example INR, OMR, SAR). Rate = how many AED one unit of the currency is worth."
fx["A17"].font = Font(name=F, size=9, italic=True)

# 5) START HERE (rewrite)
if "START HERE" in wb.sheetnames:
    del wb["START HERE"]
sh = wb.create_sheet("START HERE", 0)
sh["A1"] = "Sasco quote generator: settings workbook"; sh["A1"].font = Font(name=F, size=16, bold=True, color="1F3864")
lines = [
    ("Two workbooks", "This one holds company settings, people, Incoterm rules, terms and currencies. 'Template_design_v2' holds the line items and costs."),
    ("Settings", "Company details, quote number prefix, default currency and VAT, email wording."),
    ("FX", "Currencies, rates to AED and the currency transfer charge. Fill the yellow cells."),
    ("Sales_Team", "Name, email, role and permissions. 'Pricing + Dashboard' sees the dashboard; 'Pricing Only' does not. Active = N blocks sign-in."),
    ("Incoterm_Matrix", "Which charge categories show for each Incoterm and direction. Orange cells were changed by Claude (insurance), see column J."),
    ("Terms", "Terms and conditions by layer: common, mode, direction, load type, cargo nature, cargo type, Incoterm."),
    ("Options", "The selectable lists used by the app."),
]
for i, (a, b) in enumerate(lines, 3):
    sh.cell(row=i, column=1, value=a).font = Font(name=F, size=10, bold=True)
    c = sh.cell(row=i, column=2, value=b); c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
sh.column_dimensions["A"].width = 24; sh.column_dimensions["B"].width = 120

order = ["START HERE", "Settings", "FX", "Sales_Team", "Incoterm_Matrix", "Terms", "Options"]
wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]
wb.save(OUT)
print("matrix insurance changes:", changes)
print("email fixed:", fixed_email)
print("sheets:", wb.sheetnames)
